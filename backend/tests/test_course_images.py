"""Offline coverage for the new unillustrated-chapter image workflow."""

from io import BytesIO
import unittest
import zipfile

from docx import Document
from PIL import Image

from app.services.course_images import (
    CourseImagesError,
    Picture,
    Provider,
    _correct_existing_cover,
    build_course_image_package,
)


def jpg(width=1000, height=700, rgb=(72, 120, 150)):
    image = Image.new("RGB", (width, height), rgb)
    output = BytesIO()
    image.save(output, format="JPEG")
    return output.getvalue()


def chapter():
    doc = Document()
    doc.add_heading("Safe handling of care records", 1)
    doc.add_paragraph("Always ensure written information is secured.")
    doc.add_heading("Secure storage and access", 2)
    doc.add_paragraph("The team must securely store records and control access.")
    doc.add_heading("Data confidentiality in the workplace", 2)
    doc.add_paragraph("Staff should protect private information.")
    doc.add_heading("Responding to requests", 2)
    doc.add_paragraph("Follow the correct procedure for requests.")
    doc.add_heading("Summary", 1)
    doc.add_paragraph("The end of the chapter.")
    doc.add_heading("References", 1)
    doc.add_paragraph("Books and links.")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def pack(members):
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name, data in members:
            z.writestr(name, data)
    return out.getvalue()


class FakeProvider(Provider):
    def __init__(self):
        super().__init__(pexels_key="fake-key", openai_key="")
        self.count = 0

    def find(self, query, *, portrait=False, generate=True, prompt="", subject=""):
        self.count += 1
        return Picture(
            jpg(700, 980) if portrait else jpg(), "Pexels",
            f"https://www.pexels.com/photo/{self.count}/", str(self.count)
        )


class CourseImagesTests(unittest.TestCase):
    def test_guide_only_creates_exact_requested_folders(self):
        source = chapter()
        result = build_course_image_package(
            source, "Chapter 3.docx", course_name="Confidentiality in Care Settings",
            mode="guides", provider=FakeProvider(),
        )
        self.assertEqual(result.details["chapters_processed"], 1)
        self.assertEqual(result.details["image_recommendations"], 2)
        self.assertEqual(result.details["photos_inserted"], 0)
        self.assertEqual(result.details["covers_prepared"], 1)
        with zipfile.ZipFile(BytesIO(result.payload)) as archive:
            names = archive.namelist()
            self.assertIn("Matching Image Search Guides/3_Matching_Image_Search_Guide.docx", names)
            self.assertIn("Cover Images/Confidentiality in Care Settings_cover.png", names)
            self.assertFalse(any(name.endswith("_with_images.docx") for name in names))
            self.assertEqual(len([n for n in names if n.endswith(".docx")]), 1)
            guide = Document(BytesIO(archive.read("Matching Image Search Guides/3_Matching_Image_Search_Guide.docx")))
            all_text = " ".join(p.text for p in guide.paragraphs)
            self.assertIn("Recommended Image Placements: 2", all_text)
            self.assertNotIn("05 – Summary", all_text)
            rels = [rel.target_ref for rel in guide.part.rels.values() if rel.reltype.endswith("/hyperlink")]
            self.assertEqual(len(rels), 4)
            self.assertTrue(any("pexels.com/search/" in r for r in rels))
            self.assertTrue(any("stocksnap.io/search/" in r for r in rels))

    def test_insert_mode_preserves_all_original_text_and_styles(self):
        source = chapter()
        original = Document(BytesIO(source))
        expected = [(p.text, p.style.name) for p in original.paragraphs]
        result = build_course_image_package(
            source, "Unit 2.docx", mode="insert", course_name="Information in Care",
            provider=FakeProvider(),
        )
        self.assertEqual(result.details["photos_inserted"], 2)
        with zipfile.ZipFile(BytesIO(result.payload)) as archive:
            name = "Word Documents With Images/Unit 2_with_images.docx"
            self.assertIn(name, archive.namelist())
            doc = Document(BytesIO(archive.read(name)))
            self.assertEqual([(p.text, p.style.name) for p in doc.paragraphs if p.text], expected)
            self.assertEqual(len(doc.inline_shapes), 2)
            for label in ("Secure storage and access", "Responding to requests"):
                heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text == label)
                self.assertTrue(doc.paragraphs[heading_index + 1]._p.xpath(".//w:drawing"))
            self.assertEqual(len([p for p in doc.paragraphs if p.text == "Summary"]), 1)

    def test_does_not_offer_duplicate_image_below_existing_one(self):
        doc = Document()
        doc.add_heading("Records security", 1)
        doc.add_paragraph().add_run().add_picture(BytesIO(jpg()), width=2000000)
        doc.add_heading("Confidential information", 2)
        doc.add_paragraph("Secure the data.")
        stream = BytesIO()
        doc.save(stream)
        result = build_course_image_package(
            stream.getvalue(), "Course.docx", mode="insert", provider=FakeProvider(),
        )
        self.assertEqual(result.details["image_recommendations"], 1)
        self.assertEqual(result.details["photos_inserted"], 1)

    def test_no_api_keys_is_clearly_reported_not_faked(self):
        no_provider = Provider(pexels_key="", openai_key="")
        result = build_course_image_package(chapter(), "Lesson.docx", provider=no_provider)
        self.assertEqual(result.details["covers_prepared"], 0)
        self.assertTrue(any("cover" in w.lower() for w in result.details["warnings"]))
        with self.assertRaisesRegex(CourseImagesError, "PEXELS_API_KEY"):
            build_course_image_package(chapter(), "Lesson.docx", mode="insert", provider=no_provider)

    def test_multi_document_zip_and_excludes_assignment_templates(self):
        archive = pack([
            ("Chapter 1.docx", chapter()),
            ("nested/Chapter 2.docx", chapter()),
            ("Written Assignment Template.docx", chapter()),
            ("../outside/Chapter 12.docx", chapter()),
        ])
        result = build_course_image_package(archive, "Training.zip", provider=FakeProvider())
        self.assertEqual(result.details["chapters_processed"], 2)
        with zipfile.ZipFile(BytesIO(result.payload)) as z:
            self.assertEqual(len([n for n in z.namelist() if n.endswith("_Matching_Image_Search_Guide.docx")]), 2)

    def test_corrects_only_coloured_panel_and_preserves_photo(self):
        raw = Image.new("RGB", (900, 1250), (140, 175, 190))
        from PIL import ImageDraw
        draw = ImageDraw.Draw(raw)
        draw.rectangle((0, 930, 900, 1250), fill=(19, 150, 143))
        src = BytesIO()
        raw.save(src, format="PNG")
        corrected, note = _correct_existing_cover(src.getvalue())
        self.assertIsNotNone(corrected)
        with Image.open(BytesIO(corrected)) as image:
            self.assertEqual(image.getpixel((400, 120)), (140, 175, 190))
            self.assertEqual(image.getpixel((400, 1200)), (18, 22, 27))
            self.assertIn("charcoal", note)


if __name__ == "__main__":
    unittest.main()
