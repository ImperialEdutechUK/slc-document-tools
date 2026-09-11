import unittest
from io import BytesIO

from docx import Document

from app.formatter.simple_editing import apply_simple_edits, get_editable_paragraphs


def make_docx():
    doc = Document()
    doc.add_heading("Resources for Further Reference:", level=1)
    doc.add_paragraph("Department for Education reference one")
    doc.add_paragraph("Department for Education reference two")
    doc.add_paragraph("Body paragraph")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class SimpleEditingTests(unittest.TestCase):
    def test_lists_editable_paragraphs(self):
        paragraphs = get_editable_paragraphs(make_docx())
        texts = [item["text"] for item in paragraphs]
        self.assertIn("Resources for Further Reference:", texts)
        self.assertIn("Department for Education reference one", texts)

    def test_converts_selected_paragraphs_to_bullets(self):
        payload = make_docx()
        paragraphs = get_editable_paragraphs(payload)
        selected = [
            item["index"]
            for item in paragraphs
            if item["text"].startswith("Department for Education")
        ]

        edited, changed = apply_simple_edits(payload, selected, "bullets")
        self.assertEqual(2, changed)

        updated = get_editable_paragraphs(edited)
        updated_refs = [
            item for item in updated if item["text"].startswith("Department for Education")
        ]
        self.assertEqual(["bullet", "bullet"], [item["list_type"] for item in updated_refs])

    def test_can_force_and_remove_page_break_before(self):
        payload = make_docx()
        paragraph = next(
            item for item in get_editable_paragraphs(payload) if item["text"] == "Body paragraph"
        )
        edited, _ = apply_simple_edits(payload, [paragraph["index"]], "page_break_before")
        updated = next(
            item for item in get_editable_paragraphs(edited) if item["text"] == "Body paragraph"
        )
        self.assertTrue(updated["page_break_before"])

        normal_flow, _ = apply_simple_edits(
            edited, [updated["index"]], "remove_page_break_before"
        )
        final = next(
            item for item in get_editable_paragraphs(normal_flow) if item["text"] == "Body paragraph"
        )
        self.assertFalse(final["page_break_before"])


if __name__ == "__main__":
    unittest.main()


class ExpandedEditingTests(unittest.TestCase):
    def test_can_edit_text_and_keep_manual_new_line(self):
        from app.formatter.simple_editing import replace_paragraph_text

        payload = make_docx()
        paragraph = next(
            item
            for item in get_editable_paragraphs(payload)
            if item["text"] == "Body paragraph"
        )
        edited, changed = replace_paragraph_text(
            payload,
            paragraph["index"],
            "Updated first line\nUpdated second line",
        )
        self.assertEqual(1, changed)

        updated = next(
            item
            for item in get_editable_paragraphs(edited)
            if item["text"].startswith("Updated first line")
        )
        self.assertEqual("Updated first line\nUpdated second line", updated["text"])

    def test_can_add_blank_line_and_toggle_keep_together(self):
        from docx import Document as WordDocument

        payload = make_docx()
        paragraph = next(
            item
            for item in get_editable_paragraphs(payload)
            if item["text"] == "Body paragraph"
        )

        spaced, _ = apply_simple_edits(
            payload, [paragraph["index"]], "blank_line_before"
        )
        doc = WordDocument(BytesIO(spaced))
        texts = [p.text for p in doc.paragraphs]
        body_index = texts.index("Body paragraph")
        self.assertGreater(body_index, 0)
        self.assertEqual("", texts[body_index - 1])

        refreshed = next(
            item
            for item in get_editable_paragraphs(spaced)
            if item["text"] == "Body paragraph"
        )
        kept, _ = apply_simple_edits(
            spaced, [refreshed["index"]], "keep_together"
        )
        state = next(
            item
            for item in get_editable_paragraphs(kept)
            if item["text"] == "Body paragraph"
        )
        self.assertTrue(state["keep_together"])

        allowed, _ = apply_simple_edits(
            kept, [state["index"]], "allow_split"
        )
        final = next(
            item
            for item in get_editable_paragraphs(allowed)
            if item["text"] == "Body paragraph"
        )
        self.assertFalse(final["keep_together"])

    def test_can_read_and_edit_generated_style_footer(self):
        from docx import Document as WordDocument
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from app.formatter.simple_editing import edit_footer, get_footer_settings

        doc = WordDocument()
        doc.add_paragraph("Body")
        footer = doc.sections[0].footer
        table = footer.add_table(rows=1, cols=3, width=doc.sections[0].page_width)
        table.cell(0, 0).text = "© South London College Ltd"
        table.cell(0, 1).text = "Old Course Name"

        page_paragraph = table.cell(0, 2).paragraphs[0]
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run = page_paragraph.add_run()
        run._r.extend([begin, instr, separate])
        page_paragraph.add_run("1")
        page_paragraph.add_run(" | Page")
        run_end = page_paragraph.add_run()
        run_end._r.append(end)

        buffer = BytesIO()
        doc.save(buffer)
        payload = buffer.getvalue()

        settings = get_footer_settings(payload)
        self.assertEqual("Old Course Name", settings["course_text"])
        self.assertEqual("© South London College Ltd", settings["copyright_text"])
        self.assertEqual(" | Page", settings["page_label"])

        edited, changed_parts = edit_footer(
            payload,
            course_text="New Course Name",
            copyright_text="© SLC Test",
            page_label=" | Pg",
        )
        self.assertGreaterEqual(changed_parts, 1)
        updated = get_footer_settings(edited)
        self.assertEqual("New Course Name", updated["course_text"])
        self.assertEqual("© SLC Test", updated["copyright_text"])
        self.assertEqual(" | Pg", updated["page_label"])

    def test_can_regenerate_toc_from_current_heading_text(self):
        import zipfile
        import xml.etree.ElementTree as ET
        from app.formatter.document_structure import make_toc_elements
        from app.formatter.simple_editing import regenerate_toc, wt

        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("Some body text")
        doc.add_heading("Second Topic", level=2)
        buffer = BytesIO()
        doc.save(buffer)

        source = zipfile.ZipFile(BytesIO(buffer.getvalue()), "r")
        root = ET.fromstring(source.read("word/document.xml"))
        body = root.find(wt("body"))
        for offset, element in enumerate(make_toc_elements([])):
            body.insert(offset, element)

        document_xml = ET.tostring(
            root, encoding="UTF-8", xml_declaration=True, short_empty_elements=True
        )
        with BytesIO() as rebuilt_buffer:
            with zipfile.ZipFile(rebuilt_buffer, "w") as rebuilt:
                for info in source.infolist():
                    rebuilt.writestr(
                        info,
                        document_xml
                        if info.filename == "word/document.xml"
                        else source.read(info.filename),
                    )
            source.close()
            with_toc = rebuilt_buffer.getvalue()

        updated, entries = regenerate_toc(with_toc)
        self.assertEqual(2, entries)
        with zipfile.ZipFile(BytesIO(updated), "r") as archive:
            updated_root = ET.fromstring(archive.read("word/document.xml"))
        text = " ".join(
            node.text or "" for node in updated_root.iter(wt("t"))
        )
        self.assertIn("Introduction", text)
        self.assertIn("Second Topic", text)
