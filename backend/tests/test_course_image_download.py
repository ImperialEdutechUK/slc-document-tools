"""HTTP-level checks: insertion should download Word directly, not a guide ZIP."""

from __future__ import annotations

from io import BytesIO
from unittest import TestCase
from unittest.mock import patch
import zipfile

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.services.course_images import CourseImagesResult


class CourseImageDownloadTests(TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        session_factory = sessionmaker(bind=self.engine)

        def test_db():
            with session_factory() as session:
                yield session

        app.dependency_overrides[get_db] = test_db
        self.client = TestClient(app)
        doc = Document()
        doc.add_heading("Professional document care", 1)
        doc.add_paragraph("Keep the original content in this paragraph.")
        self.word = BytesIO()
        doc.save(self.word)

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()
        self.engine.dispose()

    def test_docx_result_has_correct_mime_and_attachment_name(self):
        result = CourseImagesResult(
            filename="Chapter_01_with_images.docx",
            payload=self.word.getvalue(),
            details={"photos_inserted": 1, "documents_with_images": 1},
        )
        with patch("app.main.build_course_image_package", return_value=result), \
             patch("app.main.storage.put_bytes") as save:
            response = self.client.post(
                "/api/v1/course-images/prepare",
                files={"document": ("Chapter_01.docx", self.word.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"mode": "insert", "image_source": "freepik"},
            )
            self.assertEqual(200, response.status_code, response.text)
            info = response.json()
            self.assertEqual(info["output_filename"], "Chapter_01_with_images.docx")
            self.assertEqual(info["details"]["documents_with_images"], 1)
            self.assertEqual(save.call_args.args[2], "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        with patch("app.main.storage.get_bytes", return_value=self.word.getvalue()):
            downloaded = self.client.get(info["download_url"])
        self.assertEqual(200, downloaded.status_code)
        self.assertEqual(downloaded.headers["content-type"], "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertIn('filename="Chapter_01_with_images.docx"', downloaded.headers["content-disposition"])
        Document(BytesIO(downloaded.content))

    def test_zip_result_stays_zip_for_multiple_documents(self):
        package = BytesIO()
        with zipfile.ZipFile(package, "w") as archive:
            archive.writestr("Chapter_01_with_images.docx", self.word.getvalue())
            archive.writestr("Chapter_02_with_images.docx", self.word.getvalue())
        result = CourseImagesResult(
            filename="My_Course_Documents_With_Images.zip",
            payload=package.getvalue(),
            details={"photos_inserted": 2, "documents_with_images": 2},
        )
        with patch("app.main.build_course_image_package", return_value=result), \
             patch("app.main.storage.put_bytes") as save:
            response = self.client.post(
                "/api/v1/course-images/prepare",
                files={"document": ("My_Course.zip", package.getvalue(), "application/zip")},
                data={"mode": "insert"},
            )
        self.assertEqual(200, response.status_code, response.text)
        self.assertTrue(response.json()["output_filename"].endswith(".zip"))
        self.assertEqual(save.call_args.args[2], "application/zip")
