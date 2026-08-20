import unittest
from io import BytesIO
from unittest.mock import MagicMock, patch

import httpx
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

from app.services.linked_images import (
    _download_stock_resource,
    _extract_resource_id,
    download_linked_images,
    scan_linked_image_sources,
)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def document_bytes(builder):
    document = Document()
    builder(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def png_bytes():
    buffer = BytesIO()
    Image.new("RGB", (20, 12), "white").save(buffer, format="PNG")
    return buffer.getvalue()


class LinkedImageTests(unittest.TestCase):
    def test_hyperlinked_placeholder_is_detected(self):
        payload = document_bytes(
            lambda doc: add_hyperlink(
                doc.add_paragraph(),
                "Image 3",
                "https://www.freepik.com/free-photo/example_12345678.htm",
            )
        )
        sources, warnings = scan_linked_image_sources(payload)
        self.assertFalse(warnings)
        self.assertEqual(1, len(sources))
        self.assertEqual(3, sources[0].number)
        self.assertEqual("freepik", sources[0].source_type)

    def test_url_in_following_paragraph_is_associated_with_placeholder(self):
        def build(doc):
            doc.add_paragraph("Image 7")
            doc.add_paragraph("https://cdn.example.com/photo.jpg")

        sources, _ = scan_linked_image_sources(document_bytes(build))
        self.assertEqual([(7, "direct")], [(item.number, item.source_type) for item in sources])


    def test_magnific_placeholder_is_detected_as_stock_source(self):
        payload = document_bytes(
            lambda doc: add_hyperlink(
                doc.add_paragraph(),
                "Image 1",
                "https://www.magnific.com/free-photo/example_25855993.htm",
            )
        )
        sources, warnings = scan_linked_image_sources(payload)
        self.assertFalse(warnings)
        self.assertEqual(1, len(sources))
        self.assertEqual("magnific", sources[0].source_type)
        self.assertEqual(25855993, _extract_resource_id(sources[0].url))

    @patch("app.services.linked_images._download_stock_resource")
    def test_magnific_link_uses_stock_api_downloader(self, downloader):
        downloader.return_value = png_bytes()

        def build(doc):
            doc.add_paragraph(
                "Image 6 https://www.magnific.com/free-vector/business-meeting_798219.htm"
            )

        result = download_linked_images(document_bytes(build))
        downloader.assert_called_once_with(
            "https://www.magnific.com/free-vector/business-meeting_798219.htm"
        )
        self.assertEqual(1, len(result.items))
        self.assertEqual("downloaded", result.entries[0].status)
        self.assertEqual("magnific", result.entries[0].source_type)

    @patch.dict("os.environ", {"MAGNIFIC_API_KEY": "test-key"}, clear=False)
    @patch("app.services.linked_images._download_public_url")
    @patch("app.services.linked_images.httpx.Client")
    def test_stock_api_uses_png_for_vector_when_available(self, client_cls, public_download):
        client = MagicMock()
        client_cls.return_value.__enter__.return_value = client

        detail = MagicMock()
        detail.json.return_value = {
            "data": {"available_formats": {"png": {"total": 1}, "svg": {"total": 1}}}
        }
        detail.raise_for_status.return_value = None

        png_response = MagicMock()
        png_response.status_code = 200
        png_response.json.return_value = {
            "data": [{"signed_url": "https://img.freepik.com/example.png"}]
        }

        client.get.side_effect = [detail, png_response]
        public_download.return_value = png_bytes()

        payload = _download_stock_resource(
            "https://www.magnific.com/free-vector/business-meeting_798219.htm"
        )

        self.assertEqual(png_bytes(), payload)
        self.assertEqual(2, client.get.call_count)
        self.assertTrue(client.get.call_args_list[1].args[0].endswith("/798219/download/png"))
        public_download.assert_called_once_with("https://img.freepik.com/example.png")


    @patch.dict("os.environ", {"MAGNIFIC_API_KEY": "test-key"}, clear=False)
    @patch("app.services.linked_images._download_public_url")
    @patch("app.services.linked_images.httpx.Client")
    def test_legacy_resource_id_searches_catalogue_and_downloads_replacement(self, client_cls, public_download):
        client = MagicMock()
        client_cls.return_value.__enter__.return_value = client

        not_found_error = httpx.HTTPStatusError(
            "404 Not Found",
            request=httpx.Request("GET", "https://api.magnific.com/v1/resources/11552633"),
            response=httpx.Response(404),
        )

        old_detail = MagicMock()
        old_detail.raise_for_status.side_effect = not_found_error

        old_jpg = MagicMock()
        old_jpg.status_code = 404
        old_png = MagicMock()
        old_png.status_code = 404

        old_default = MagicMock()
        old_default.raise_for_status.side_effect = not_found_error

        search = MagicMock()
        search.raise_for_status.return_value = None
        search.json.return_value = {
            "data": [
                {
                    "id": 9911552633,
                    "title": "Real estate property immovables assessment business",
                    "url": "https://www.magnific.com/free-vector/real-estate-property-immovables-assessment-business_9911552633.htm",
                    "meta": {"available_formats": {"png": {"total": 1}}},
                }
            ]
        }

        new_png = MagicMock()
        new_png.status_code = 200
        new_png.json.return_value = {
            "data": [{"signed_url": "https://img.magnific.com/replacement.png"}]
        }

        client.get.side_effect = [old_detail, old_jpg, old_png, old_default, search, new_png]
        public_download.return_value = png_bytes()

        payload = _download_stock_resource(
            "https://www.magnific.com/free-vector/real-estate-property-immovables-assessment-business_11552633.htm"
        )

        self.assertEqual(png_bytes(), payload)
        search_call = client.get.call_args_list[4]
        self.assertEqual("https://api.magnific.com/v1/resources", search_call.args[0])
        self.assertIn("real estate property immovables assessment business", search_call.kwargs["params"]["term"])
        self.assertTrue(client.get.call_args_list[5].args[0].endswith("/9911552633/download/png"))

    @patch.dict("os.environ", {"MAGNIFIC_API_KEY": "test-key"}, clear=False)
    @patch("app.services.linked_images.httpx.Client")
    def test_legacy_resource_does_not_accept_unrelated_catalogue_result(self, client_cls):
        client = MagicMock()
        client_cls.return_value.__enter__.return_value = client

        not_found_error = httpx.HTTPStatusError(
            "404 Not Found",
            request=httpx.Request("GET", "https://api.magnific.com/v1/resources/11552633"),
            response=httpx.Response(404),
        )
        old_detail = MagicMock()
        old_detail.raise_for_status.side_effect = not_found_error
        old_jpg = MagicMock(status_code=404)
        old_png = MagicMock(status_code=404)
        old_default = MagicMock()
        old_default.raise_for_status.side_effect = not_found_error
        search = MagicMock()
        search.raise_for_status.return_value = None
        search.json.return_value = {
            "data": [{"id": 123, "title": "Completely unrelated beach photograph", "url": "https://www.magnific.com/free-photo/beach_123.htm"}]
        }
        client.get.side_effect = [old_detail, old_jpg, old_png, old_default, search]

        with self.assertRaisesRegex(ValueError, "no high-confidence replacement"):
            _download_stock_resource(
                "https://www.magnific.com/free-vector/real-estate-property-immovables-assessment-business_11552633.htm"
            )

    @patch("app.services.linked_images._download_public_url")
    def test_direct_link_download_is_returned_as_numbered_image(self, downloader):
        downloader.return_value = png_bytes()

        def build(doc):
            doc.add_paragraph("Image 2 https://cdn.example.com/photo.png")

        result = download_linked_images(document_bytes(build))
        self.assertEqual(1, len(result.items))
        self.assertEqual("Image 2.jpg", result.items[0][0])
        self.assertEqual("downloaded", result.entries[0].status)

    @patch("app.services.linked_images._download_public_url")
    def test_manual_image_takes_priority(self, downloader):
        def build(doc):
            doc.add_paragraph("Image 4 https://cdn.example.com/photo.png")

        result = download_linked_images(
            document_bytes(build),
            [("4.jpg", png_bytes())],
        )
        downloader.assert_not_called()
        self.assertEqual([], result.items)
        self.assertEqual("manual_override", result.entries[0].status)


if __name__ == "__main__":
    unittest.main()
