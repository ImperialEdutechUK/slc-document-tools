"""Search guides, cover images and optional in-document imagery for image-free courses.

The original course DOCX is never reformatted: insert mode only appends standalone
image paragraphs after selected headings. Guide mode never returns modified DOCX.
External image providers are opt-in via server-side environment variables.
"""

from __future__ import annotations

import base64
from dataclasses import dataclass, field
from io import BytesIO
import math
import os
from pathlib import Path, PurePosixPath
import re
from urllib.parse import quote, quote_plus, urlparse
import zipfile

import httpx
from PIL import Image, ImageDraw, ImageStat, ImageOps, UnidentifiedImageError
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from .zip_utils import normalise_zip_name, safe_zip_member

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_BYTES = 200 * 1024 * 1024
MAX_IMAGE_BYTES = 22 * 1024 * 1024
MAX_ARCHIVE_FILES = 160
SKIP_HEADING = re.compile(
    r"\b(summary|summaries|conclusions?|references?|bibliograph(?:y|ies)|"
    r"table of contents|contents|resources for further reference)\b", re.I
)
SKIP_DOC = re.compile(r"\b(assignment[\s_-]*brief|written[\s_-]*assignment[\s_-]*template)\b", re.I)
TOKEN_RE = re.compile(r"[a-zA-Z][a-zA-Z'-]+")
STOPWORDS = frozenset(
    "the a an and or for of to in on at with from by about into it their your how what "
    "course chapter unit module lesson level award diploma certificate qls rqf ncfe cache "
    "using use understanding introduction principles practice practices importance overview "
    "key types role roles study studying skills basic activity activities different and".split()
)


class CourseImagesError(ValueError):
    """User-visible bad input or missing image-provider configuration."""


@dataclass
class Chapter:
    filename: str
    payload: bytes
    number: str
    title: str


@dataclass
class ImagePlan:
    paragraph_index: int
    heading: str
    search: str
    description: str
    location: str
    context: str = ""
    source: str = ""
    source_url: str = ""

    @property
    def pexels_link(self) -> str:
        return "https://www.pexels.com/search/" + quote(self.search.strip(), safe="") + "/"

    @property
    def stocksnap_link(self) -> str:
        return "https://stocksnap.io/search/" + quote_plus(self.search.strip())


@dataclass
class Picture:
    payload: bytes
    source: str
    source_url: str = ""
    identifier: str = ""


@dataclass
class CourseImagesResult:
    filename: str
    payload: bytes
    details: dict


@dataclass
class Provider:
    """Try a relevant free stock photograph, then optionally generated imagery."""

    pexels_key: str = field(default_factory=lambda: os.getenv("PEXELS_API_KEY", "").strip())
    openai_key: str = field(default_factory=lambda: os.getenv("OPENAI_API_KEY", "").strip())
    image_model: str = field(default_factory=lambda: os.getenv("OPENAI_IMAGE_MODEL", "gpt-image-2.5-flare"))
    used_photos: set[str] = field(default_factory=set)
    warnings: list[str] = field(default_factory=list)

    @property
    def available(self) -> bool:
        return bool(self.pexels_key or self.openai_key)

    def find(
        self, query: str, *, portrait: bool = False, generate: bool = True,
        prompt: str = "", subject: str = "",
    ) -> Picture | None:
        if self.pexels_key:
            try:
                found = self._stock(query, portrait=portrait, subject=subject)
                if found:
                    return found
            except (httpx.HTTPError, ValueError, OSError) as exc:
                self.warnings.append(f"Pexels lookup was unsuccessful: {type(exc).__name__}.")
        if generate and self.openai_key:
            try:
                return self._generate(prompt or query, portrait=portrait)
            except (httpx.HTTPError, ValueError, OSError, KeyError) as exc:
                self.warnings.append(f"AI image generation was unsuccessful: {type(exc).__name__}.")
        return None

    def _stock(self, query: str, *, portrait: bool, subject: str = "") -> Picture | None:
        headers = {"Authorization": self.pexels_key}
        with httpx.Client(timeout=22, follow_redirects=False) as client:
            response = client.get(
                "https://api.pexels.com/v1/search",
                params={"query": query[:110], "orientation": "portrait" if portrait else "landscape", "per_page": 20},
                headers=headers,
            )
            response.raise_for_status()
            result = response.json()
            photos = result.get("photos") or []
            tokens = set(_tokens(query))
            # For automatic insertion, a search-result caption must mention a
            # meaningful term from the requested heading/cover subject, not
            # merely generic terms such as 'workplace' or 'professional'.
            topic_tokens = set(_tokens(subject)) - {"workplace", "professional", "course", "setting", "settings"}
            ranked: list[tuple[int, dict]] = []
            for photo in photos:
                identifier = str(photo.get("id", ""))
                if not identifier or identifier in self.used_photos:
                    continue
                alt = set(_tokens(str(photo.get("alt", ""))))
                # Don't silently accept unrelated search hits. A relevant
                # subject token must be represented in Pexels' photo caption.
                overlap = len(tokens & alt)
                if overlap < 1 or (topic_tokens and not (topic_tokens & alt)):
                    continue
                ranked.append((overlap, photo))
            ranked.sort(key=lambda pair: pair[0], reverse=True)
            for _, photo in ranked[:5]:
                src = photo.get("src") or {}
                url = src.get("large2x") or src.get("large") or src.get("original")
                if not url or urlparse(url).scheme != "https" or urlparse(url).hostname != "images.pexels.com":
                    continue
                try:
                    with client.stream("GET", url) as stream:
                        stream.raise_for_status()
                        chunks: list[bytes] = []
                        total = 0
                        for chunk in stream.iter_bytes():
                            total += len(chunk)
                            if total > MAX_IMAGE_BYTES:
                                raise ValueError("Image exceeds the upload limit")
                            chunks.append(chunk)
                    payload = _normalised_image(b"".join(chunks), portrait=portrait)
                except (httpx.HTTPError, ValueError, OSError, UnidentifiedImageError):
                    continue
                identifier = str(photo["id"])
                self.used_photos.add(identifier)
                return Picture(payload, "Pexels", str(photo.get("url", "")), identifier)
        return None

    def _generate(self, prompt: str, *, portrait: bool) -> Picture:
        request = {
            "model": self.image_model,
            "prompt": (
                "Professional, candid, ultra-photorealistic educational stock photograph. "
                "Authentic setting, realistic people and hands, natural lighting, credible "
                "equipment and details, visually distinct photographic composition. "
                "No text, writing, logos, watermarks, diagrams, synthetic-looking faces or stylisation. "
                + prompt[:2200]
            ),
            "size": "1024x1536" if portrait else "1536x1024",
            "quality": "medium",
            "output_format": "png",
            "n": 1,
        }
        with httpx.Client(timeout=180) as client:
            response = client.post(
                "https://api.openai.com/v1/images/generations",
                headers={"Authorization": f"Bearer {self.openai_key}"},
                json=request,
            )
            response.raise_for_status()
            items = response.json().get("data") or []
            if not items or not items[0].get("b64_json"):
                raise ValueError("Image API returned no image")
            encoded = items[0]["b64_json"]
            if len(encoded) > 4 * MAX_IMAGE_BYTES // 3 + 100:
                raise ValueError("Image API payload too large")
            payload = base64.b64decode(encoded, validate=True)
            return Picture(_normalised_image(payload, portrait=portrait), "AI-generated")


def _normalised_image(raw: bytes, *, portrait: bool) -> bytes:
    if not raw or len(raw) > MAX_IMAGE_BYTES:
        raise ValueError("Invalid or oversized image")
    with Image.open(BytesIO(raw)) as source:
        source.load()
        if source.width < 450 or source.height < 300:
            raise ValueError("Image resolution too small")
        if source.width * source.height > 65_000_000:
            raise ValueError("Image resolution too large")
        image = ImageOps.exif_transpose(source).convert("RGB")
        if portrait:
            image.thumbnail((1450, 2100), Image.Resampling.LANCZOS)
        else:
            # Visual assets below headings should always be landscape.
            image = ImageOps.fit(image, (1300, 812), method=Image.Resampling.LANCZOS)
        target = BytesIO()
        image.save(target, format="JPEG", quality=89, optimize=True)
        return target.getvalue()


def _tokens(value: str) -> list[str]:
    return [t.lower().strip("-'") for t in TOKEN_RE.findall(value) if len(t) > 2 and t.lower() not in STOPWORDS]


def _safe_name(filename: str) -> str:
    return re.sub(r"[^a-zA-Z0-9 _.-]", "_", filename).strip(" .")[:120] or "Course"


def _number(filename: str, index: int) -> str:
    stem = Path(filename).stem
    match = re.search(r"(?:chapter|unit|module)\s*(\d+[a-z]?)\b", stem, re.I) or re.match(r"^(\d+[a-z]?)\b", stem, re.I)
    return match.group(1) if match else f"{index:02d}"


def _chapter_title(filename: str) -> str:
    return " ".join(re.sub(r"[_-]+", " ", Path(filename).stem).split())


def _read_chapters(raw: bytes, filename: str) -> tuple[list[Chapter], list[tuple[str, bytes]]]:
    ext = Path(filename).suffix.lower()
    if ext == ".docx":
        if len(raw) > MAX_DOCUMENT_BYTES:
            raise CourseImagesError("DOCX is too large (25 MB maximum).")
        return [Chapter(filename, raw, _number(filename, 1), _chapter_title(filename))], []
    if ext != ".zip":
        raise CourseImagesError("Upload a DOCX document or a ZIP containing Word documents.")
    chapters: list[Chapter] = []
    covers: list[tuple[str, bytes]] = []
    total_size = 0
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if len(archive.infolist()) > MAX_ARCHIVE_FILES:
                raise CourseImagesError("Archive contains too many entries.")
            for info in archive.infolist():
                if info.is_dir() or not safe_zip_member(info.filename):
                    continue
                name = PurePosixPath(normalise_zip_name(info.filename)).name
                extension = Path(name).suffix.lower()
                if extension not in {".docx", ".png", ".jpg", ".jpeg", ".webp"}:
                    continue
                if info.file_size > MAX_DOCUMENT_BYTES:
                    raise CourseImagesError(f"Archive entry is too large: {name}")
                total_size += info.file_size
                if total_size > MAX_ARCHIVE_BYTES:
                    raise CourseImagesError("Archive contains too much uncompressed data.")
                if extension == ".docx" and not SKIP_DOC.search(_chapter_title(name)):
                    chapters.append(Chapter(name, archive.read(info), _number(name, len(chapters) + 1), _chapter_title(name)))
                elif extension != ".docx" and re.search(r"\bcover\b", re.sub(r"[_-]", " ", Path(name).stem), re.I):
                    covers.append((name, archive.read(info)))
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise CourseImagesError("The uploaded ZIP is not readable.") from exc
    if not chapters:
        raise CourseImagesError("No chapter DOCX files were found in the uploaded archive.")
    return chapters, covers


def _picture_after_heading(paragraphs, index: int) -> bool:
    # Check the paragraph directly after the heading and a following spacer,
    # but never a later section or unrelated paragraph.
    for p in paragraphs[index + 1:index + 3]:
        if p.style and re.match(r"^Heading\s*\d", p.style.name or "", re.I):
            break
        if p._p.xpath(".//w:drawing") or p._p.xpath(".//w:pict"):
            return True
        if p.text.strip():
            break
    return False


def _scene_context(course: str, heading: str, excerpt: str) -> tuple[str, str]:
    combined = (course + " " + heading).lower()
    if any(w in combined for w in ("confidential", "information", "data protection", "records")):
        setting = "care or office information-handling workplace" if "care" in combined else "professional office"
        action = "a professional working carefully with secure records or workplace information"
        stock_tags = "secure records workplace"
    elif any(w in combined for w in ("food", "hygiene", "catering", "kitchen")):
        setting = "professional kitchen or food-preparation area"
        action = "a catering or care professional demonstrating a relevant food-safety activity"
        stock_tags = "professional food safety kitchen"
    elif any(w in combined for w in ("care", "dementia", "health", "patient", "nurs")):
        setting = "authentic health or social care setting"
        action = "a professional supporting a person or demonstrating a relevant care activity"
        stock_tags = "healthcare staff care"
    elif any(w in combined for w in ("child", "teaching", "education", "school", "learning")):
        setting = "authentic classroom or early-years learning setting"
        action = "educators or learners participating in a relevant activity"
        stock_tags = "education classroom teacher"
    elif any(w in combined for w in ("management", "leadership", "business", "finance", "customer")):
        setting = "credible UK business or workplace environment"
        action = "professionals undertaking a relevant everyday business activity"
        stock_tags = "business office meeting"
    else:
        setting = "realistic workplace appropriate to the course topic"
        action = "people, activities or equipment demonstrating the specific heading"
        stock_tags = "professional work"
    # Tie each shot to the heading's meaning, not just the course. The
    # surrounding chapter text is used as context, never altered or invented.
    heading_lower = heading.lower()
    scenes = (
        (r"storage|storing|filing|access|records?", "a staff member placing a confidential folder in a locked cabinet beside a secure workstation"),
        (r"requests?|responding|enquiries|questions?", "two staff discussing an information request discreetly at a workplace reception desk"),
        (r"consent|permission|choices?", "a care professional explaining paperwork to an adult in a private consultation room"),
        (r"confidential|privacy|data protect|security|password", "a worker protecting private records at a secure computer station with no readable personal information"),
        (r"handwashing|washing|hygiene|infection", "a professional using a handwashing sink and appropriate hygiene supplies"),
        (r"cook|temperature|heating|chill|storage of food", "a trained kitchen worker checking food with an appropriate digital probe thermometer"),
        (r"risk|hazard|safety check|assess", "a workplace supervisor checking relevant equipment against a paper safety checklist"),
        (r"equality|diversity|inclusion", "colleagues of different backgrounds working collaboratively in a respectful professional setting"),
        (r"leadership|teamwork|team management|delegat", "a manager coordinating colleagues during a realistic planning discussion"),
        (r"review|audit|monitor|evaluat", "a supervisor reviewing an activity checklist alongside a colleague"),
        (r"training|teach|learn|demonstrat", "a professional demonstrating a relevant practical skill to a colleague"),
    )
    for pattern, specific_action in scenes:
        if re.search(pattern, heading_lower):
            action = specific_action
            break
    topic_tokens = _tokens(heading)[:5] or _tokens(excerpt)[:3]
    subject_tokens = _tokens(course)[:3]
    query = " ".join(dict.fromkeys(topic_tokens + subject_tokens + _tokens(stock_tags)[:3]))[:105]
    if not query:
        query = stock_tags
    description = (
        f"An ultra-realistic landscape photograph for the section '{heading}' "
        f"in a {setting}, showing {action}. Use believable people, activities and equipment "
        "in a natural environment, with no promotional text, watermarks or legible private data."
    )
    return query, description


def _plans(document: Document, course_title: str) -> tuple[list[ImagePlan], int]:
    paragraphs = document.paragraphs
    eligible: list[ImagePlan] = []
    for i, p in enumerate(paragraphs):
        text = " ".join(p.text.split())
        if not text or SKIP_HEADING.search(text):
            continue
        if not p.style or not re.match(r"^Heading\s*[1-6]$", p.style.name or "", re.I):
            continue
        if _picture_after_heading(paragraphs, i):
            continue
        context_items = []
        for following in paragraphs[i + 1:i + 7]:
            if following.style and re.match(r"^Heading\s*[1-6]$", following.style.name or "", re.I):
                break
            if following.text.strip():
                context_items.append(" ".join(following.text.split()))
            if sum(len(x) for x in context_items) >= 380:
                break
        excerpt = " ".join(context_items)[:450]
        # Heading-only chapter dividers have no instructional content under
        # them and should not receive filler or decorative photography.
        if not excerpt:
            continue
        query, description = _scene_context(course_title, text, excerpt)
        eligible.append(ImagePlan(i, text, query, description, f"Immediately below the heading '{text}' in the original Word document.", excerpt))
    count = math.ceil(len(eligible) * 0.5)
    if not count:
        return [], len(eligible)
    selected = [eligible[min(len(eligible) - 1, int((i + 0.5) * len(eligible) / count))] for i in range(count)]
    return selected, len(eligible)


def _hyperlink(paragraph: Paragraph, label: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "146C89")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _write_guide(course_title: str, chapter: Chapter, plans: list[ImagePlan]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(18)
    heading_run = title.add_run("Matching Image Search Guide")
    heading_run.bold = True
    heading_run.font.size = Pt(22)
    heading_run.font.color.rgb = RGBColor(18, 88, 108)
    for label, value in (
        ("Course / Module Title", course_title),
        ("Chapter", f"{chapter.number} – {chapter.title}"),
        ("Recommended Image Placements", str(len(plans))),
    ):
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(value)
    if not plans:
        doc.add_paragraph("No eligible headings without images were identified in this chapter.")
    for index, plan in enumerate(plans, start=1):
        h = doc.add_heading(f"{index:02d} – {plan.heading}", level=2)
        h.paragraph_format.keep_with_next = True
        for label, value in (("Photo", plan.description), ("Place", plan.location)):
            p = doc.add_paragraph()
            p.add_run(label + ": ").bold = True
            p.add_run(value)
        p = doc.add_paragraph()
        p.add_run("Links: ").bold = True
        _hyperlink(p, "Pexels – search photos", plan.pexels_link)
        p = doc.add_paragraph()
        _hyperlink(p, "StockSnap – search photos", plan.stocksnap_link)
        if plan.source_url:
            p = doc.add_paragraph()
            p.add_run("Selected photograph: ").bold = True
            _hyperlink(p, plan.source, plan.source_url)
        elif plan.source:
            p = doc.add_paragraph()
            p.add_run("Image used: ").bold = True
            p.add_run(plan.source)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _insert_images(document: Document, plans: list[ImagePlan], images: dict[int, Picture]) -> bytes:
    # Snapshot the original paragraphs first. document.paragraphs is dynamic;
    # inserted picture paragraphs would otherwise shift later heading indices.
    original_paragraphs = list(document.paragraphs)
    for plan in plans:
        photo = images.get(plan.paragraph_index)
        if not photo:
            continue
        source = original_paragraphs[plan.paragraph_index]
        new_xml = OxmlElement("w:p")
        source._p.addnext(new_xml)
        p = Paragraph(new_xml, source._parent)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.keep_together = True
        # Use the actual page's text width. Avoid changing margins or styles.
        section = document.sections[0]
        max_width = section.page_width - section.left_margin - section.right_margin
        width_inches = min(5.8, max_width / Inches(1))
        p.add_run().add_picture(BytesIO(photo.payload), width=Inches(width_inches))
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _compose_cover(picture: bytes) -> bytes:
    with Image.open(BytesIO(picture)) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        canvas = Image.new("RGB", (1240, 1754), (18, 22, 27))
        top = 1280  # 73% photographic area, 27% charcoal title panel.
        visual = ImageOps.fit(image, (1240, top + 2), method=Image.Resampling.LANCZOS)
        canvas.paste(visual, (0, 0))
        gradient = Image.new("RGBA", (1240, 90), (0, 0, 0, 0))
        mask = Image.new("L", (1, 90))
        for y in range(90):
            mask.putpixel((0, y), int((y / 89) ** 1.6 * 95))
        overlay = Image.new("RGBA", (1240, 90), (0, 0, 0, 220))
        canvas.paste(overlay, (0, top - 90), mask.resize((1240, 90)))
        draw = ImageDraw.Draw(canvas)
        draw.rectangle((0, top - 2, 1240, top + 3), fill=(196, 162, 77))
        output = BytesIO()
        canvas.save(output, format="PNG", optimize=True)
        return output.getvalue()


def _correct_existing_cover(picture: bytes) -> tuple[bytes | None, str]:
    """Correct a solid lower coloured title panel without redrawing the photo."""
    try:
        with Image.open(BytesIO(picture)) as original:
            image = ImageOps.exif_transpose(original).convert("RGB")
            width, height = image.size
            if width < 400 or height < 550 or height <= width:
                return None, "Cover was not portrait; existing photo was not altered."
            bottom = image.crop((int(width * .13), int(height * .86), int(width * .87), int(height * .95)))
            stat = ImageStat.Stat(bottom)
            if max(stat.mean) < 86 and max(stat.stddev) < 35:
                return picture, "Existing charcoal cover retained unchanged."
            # Only edit covers where there is clearly a reasonably uniform
            # lower colour panel. Otherwise subjects/photo might be obscured.
            if max(stat.stddev) > 27:
                return None, "Existing cover has no reliably detected solid lower panel; left unchanged."
            panel_start = int(height * .74)
            # Find a level above the sampled panel that remains similarly flat
            # in colour. This keeps every pixel of the original photo intact.
            bottom_colour = stat.mean
            for y in range(int(height * .61), int(height * .83), max(1, height // 100)):
                sample = ImageStat.Stat(image.crop((int(width * .15), y, int(width * .85), min(height, y + max(4, height // 80)))))
                if max(sample.stddev) < 28 and sum(abs(a - b) for a, b in zip(sample.mean, bottom_colour)) < 55:
                    panel_start = y
                    break
            draw = ImageDraw.Draw(image)
            draw.rectangle((0, panel_start, width, height), fill=(18, 22, 27))
            draw.rectangle((0, panel_start - 2, width, panel_start + 2), fill=(196, 162, 77))
            output = BytesIO()
            image.save(output, format="PNG", optimize=True)
            return output.getvalue(), "Existing coloured panel changed to charcoal; upper photograph preserved."
    except (OSError, UnidentifiedImageError):
        return None, "Existing cover image could not be read."


def _cover_for_course(
    course_title: str, supplied: tuple[str, bytes] | None, provider: Provider, *, generate: bool
) -> tuple[bytes | None, str]:
    if supplied:
        corrected, description = _correct_existing_cover(supplied[1])
        if corrected:
            return corrected, description
        # Don't crop or redraw a supplied cover whose lower panel cannot
        # safely be identified: that could erase a person or existing artwork.
        return None, description
    search = " ".join(_tokens(course_title)[:6]) + " workplace professional"
    photo = provider.find(
        search.strip(), portrait=True, generate=generate, subject=course_title,
        prompt=f"Portrait cover photograph for UK training course titled '{course_title}'. "
        "Show a directly relevant professional scene. Main subjects in upper three quarters "
        "with a natural photographic composition, no titles or graphics.",
    )
    if not photo:
        return None, "No relevant cover photograph was available; cover not fabricated."
    return _compose_cover(photo.payload), f"Cover photograph: {photo.source}" + (f" ({photo.source_url})" if photo.source_url else "")


def build_course_image_package(
    raw: bytes, filename: str, *, course_name: str = "", mode: str = "guides",
    generate_missing: bool = True, cover_upload: bytes | None = None,
    provider: Provider | None = None,
) -> CourseImagesResult:
    """Produce exact guide+cover folders or optional non-destructive DOCX enhancement."""
    if mode not in {"guides", "insert"}:
        raise CourseImagesError("Choose search guides or automatic image insertion.")
    if not raw:
        raise CourseImagesError("The uploaded document is empty.")
    chapters, provided_covers = _read_chapters(raw, filename)
    if cover_upload:
        if len(cover_upload) > MAX_IMAGE_BYTES:
            raise CourseImagesError("Cover upload is too large.")
        provided_covers = [("uploaded_cover.png", cover_upload)]
    provider = provider if provider is not None else Provider()
    if mode == "insert" and not provider.available:
        raise CourseImagesError(
            "Automatic insertion needs a PEXELS_API_KEY or OPENAI_API_KEY configured "
            "on the backend. Search guides can be created without either key."
        )

    enhanced: list[tuple[str, bytes]] = []
    guides: list[tuple[str, bytes]] = []
    cover_assets: list[tuple[str, bytes]] = []
    warnings: list[str] = []
    photos_inserted = 0
    total_recommended = 0
    used_guide_names: set[str] = set()
    course_title = course_name.strip() or (
        _chapter_title(filename) if Path(filename).suffix.lower() == ".zip" else chapters[0].title
    )
    for chapter in chapters:
        try:
            document = Document(BytesIO(chapter.payload))
            plans, eligible_count = _plans(document, course_title + " " + chapter.title)
            total_recommended += len(plans)
            acquired: dict[int, Picture] = {}
            if mode == "insert":
                for plan in plans:
                    photo = provider.find(
                        plan.search, generate=generate_missing, subject=plan.heading,
                        prompt=f"UK educational chapter '{chapter.title}', course '{course_title}', "
                        f"heading '{plan.heading}'. {plan.description}. "
                        f"Supporting passage: {plan.context[:340]}",
                    )
                    if photo:
                        acquired[plan.paragraph_index] = photo
                        plan.source = photo.source
                        plan.source_url = photo.source_url
                    else:
                        warnings.append(f"{chapter.filename}: no suitable image found for '{plan.heading}'.")
                photos_inserted += len(acquired)
                enhanced.append((f"{Path(chapter.filename).stem}_with_images.docx", _insert_images(document, plans, acquired)))
            name = f"{chapter.number}_Matching_Image_Search_Guide.docx"
            if name.lower() in used_guide_names:
                name = f"{chapter.number}_{len(guides)+1}_Matching_Image_Search_Guide.docx"
            used_guide_names.add(name.lower())
            guides.append((name, _write_guide(course_title, chapter, plans)))
            if not plans:
                warnings.append(f"{chapter.filename}: no eligible headings without an existing image were found.")
        except Exception as exc:
            # Keep all other chapters progressing without claiming this one
            # succeeded. Includes malformed DOCX payloads.
            warnings.append(f"{chapter.filename}: could not be processed ({type(exc).__name__}).")
            if mode == "insert":
                enhanced.append((f"{Path(chapter.filename).stem}_unchanged.docx", chapter.payload))

    # A shared course cover is preferable when there are many chapters of one
    # course; otherwise prepare one per separately titled chapter.
    # In a multi-chapter course with no chapter-specific covers, create one
    # shared course cover rather than repeating the same photo for each unit.
    cover_targets = (
        [course_title]
        if course_name.strip() or len(chapters) == 1 or len(provided_covers) <= 1
        else [chapter.title for chapter in chapters]
    )
    for i, title in enumerate(cover_targets):
        supplied = None
        if provided_covers:
            if len(provided_covers) == 1:
                supplied = provided_covers[0]
            else:
                supplied = next((item for item in provided_covers if _tokens(title)[-1:] and _tokens(title)[-1] in item[0].lower()), None)
        cover, note = _cover_for_course(title, supplied, provider, generate=generate_missing)
        if cover:
            cover_assets.append((f"{_safe_name(title)}_cover.png", cover))
        else:
            warnings.append(f"{title}: {note}")

    if not guides:
        raise CourseImagesError("None of the chapter documents could be read as Word files.")
    out = BytesIO()
    used_paths: set[str] = set()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        names = ("Matching Image Search Guides", "Cover Images")
        if mode == "insert":
            names += ("Word Documents With Images",)
        for folder in names:
            archive.writestr(folder + "/", b"")
        for folder, items in (
            ("Matching Image Search Guides", guides),
            ("Cover Images", cover_assets),
            ("Word Documents With Images", enhanced),
        ):
            for original_name, data in items:
                name = _safe_name(original_name)
                path = folder + "/" + name
                number = 2
                while path.lower() in used_paths:
                    path = folder + "/" + f"{Path(name).stem}_{number}{Path(name).suffix}"
                    number += 1
                used_paths.add(path.lower())
                archive.writestr(path, data)
    warnings.extend(provider.warnings)
    report = {
        "chapters_processed": len(guides),
        "image_recommendations": total_recommended,
        "photos_inserted": photos_inserted,
        "covers_prepared": len(cover_assets),
        "mode": mode,
        "warnings": list(dict.fromkeys(warnings)),
        "providers_configured": {"pexels": bool(provider.pexels_key), "ai_generation": bool(provider.openai_key)},
    }
    return CourseImagesResult(
        filename=f"{_safe_name(course_title)}_Matching_Image_Guides_and_Covers.zip" if mode == "guides"
        else f"{_safe_name(course_title)}_Documents_With_Images.zip",
        payload=out.getvalue(),
        details=report,
    )
